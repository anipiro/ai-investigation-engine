import io
import json

import pytest

from engine.ingestion.chunking import chunk_fixed, chunk_recursive, chunk_text
from engine.ingestion.dedup import content_hash, hamming, simhash64
from engine.ingestion.parsers import ParseError, UnsupportedFormatError, parse_file


# ---------------------------------------------------------------- parsers
def test_parse_markdown_title_and_whitespace():
    doc = parse_file(b"# Incident 42\r\n\r\n\r\n\r\nSomething   broke.", "notes.md")
    assert doc.title == "Incident 42"
    assert doc.fmt == "markdown"
    assert "\n\n\n" not in doc.text and "Something broke." in doc.text


def test_parse_html_strips_scripts_and_keeps_headings():
    html = b"<html><head><title>Change log</title><script>alert(1)</script></head><body><h2>CHG-1</h2><p>Pool 200 to 80</p></body></html>"
    doc = parse_file(html, "c.html")
    assert doc.title == "Change log"
    assert "alert" not in doc.text
    assert "## CHG-1" in doc.text and "Pool 200 to 80" in doc.text


def test_parse_json_records():
    data = json.dumps([{"time": "10:15", "text": "Investigating"}, {"time": "11:00", "text": "Resolved"}]).encode()
    doc = parse_file(data, "status.json")
    assert doc.metadata["records"] == 2
    assert "time: 10:15" in doc.text and "Resolved" in doc.text


def test_parse_docx_headings():
    import docx

    d = docx.Document()
    d.core_properties.title = "Postmortem"
    d.add_heading("Root cause", level=1)
    d.add_paragraph("Pool exhaustion.")
    buf = io.BytesIO()
    d.save(buf)
    doc = parse_file(buf.getvalue(), "pm.docx")
    assert doc.title == "Postmortem" and "# Root cause" in doc.text and "Pool exhaustion." in doc.text


def test_parse_pdf_pages():
    from fpdf import FPDF

    pdf = FPDF()
    for text in ("first page text", "second page text"):
        pdf.add_page()
        pdf.set_font("Helvetica", size=12)
        pdf.multi_cell(0, 8, text)
    doc = parse_file(bytes(pdf.output()), "r.pdf")
    assert doc.metadata["pages"] == 2
    assert "\f" in doc.text and "second page text" in doc.text


def test_unsupported_and_empty():
    with pytest.raises(UnsupportedFormatError):
        parse_file(b"x", "a.exe")
    with pytest.raises(ParseError):
        parse_file(b"   ", "a.txt")
    with pytest.raises(ParseError):
        parse_file(b"{not json", "a.json")


# ---------------------------------------------------------------- chunking
LONG = " ".join(f"Sentence number {i} talks about the ledger pool." for i in range(120))


def test_fixed_chunks_respect_size_and_overlap():
    chunks = chunk_fixed(LONG, max_words=50, overlap=10)
    assert all(c.n_words <= 50 for c in chunks)
    first, second = chunks[0].text.split(), chunks[1].text.split()
    assert first[-10:] == second[:10]


def test_recursive_never_splits_sentences_and_tracks_sections():
    text = "# Report\n\n## Impact\n\n" + LONG + "\n\n## Cause\n\nThe pool was reduced."
    chunks = chunk_recursive(text, max_words=60, overlap=15)
    assert all(c.n_words <= 60 for c in chunks)
    assert all(c.text.endswith(".") for c in chunks)
    assert chunks[0].section == "Report > Impact"
    assert chunks[-1].section == "Report > Cause" and chunks[-1].text == "The pool was reduced."


def test_page_numbers_from_form_feeds():
    chunks = chunk_text("page one text.\n\f\npage two text.", "recursive")
    assert [c.page for c in chunks] == [1, 2]


def test_unknown_strategy():
    with pytest.raises(ValueError):
        chunk_text("x", "semantic-magic")


# ---------------------------------------------------------------- dedup
def test_content_hash_normalises():
    assert content_hash("Hello,   World!") == content_hash("hello world")
    assert content_hash("hello world") != content_hash("hello there world")


def test_simhash_near_duplicates():
    a = "The ledger-db pool was reduced from 200 to 80 connections during the Q1 cost review. " * 5
    b = a.replace("Q1", "first quarter", 1)
    c = "Completely different text about film directors and their nationalities. " * 5
    assert hamming(simhash64(a), simhash64(b)) <= 3
    assert hamming(simhash64(a), simhash64(c)) > 10
