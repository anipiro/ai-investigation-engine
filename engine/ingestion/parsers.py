"""File-format parsers: bytes -> ParsedDocument (plain text + metadata).

Supported: .txt, .md, .pdf, .docx, .html/.htm, .json, .jsonl
Each parser is a small pure function so new formats can be registered in PARSERS.
"""

from __future__ import annotations

import io
import json
import re
from dataclasses import dataclass, field
from pathlib import PurePath
from typing import Callable


class UnsupportedFormatError(ValueError):
    pass


class ParseError(ValueError):
    pass


@dataclass
class ParsedDocument:
    text: str
    title: str
    fmt: str
    metadata: dict = field(default_factory=dict)


_WS = re.compile(r"[ \t ]+")
_BLANKS = re.compile(r"\n{3,}")


def normalize_whitespace(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = "\n".join(_WS.sub(" ", line).strip() for line in text.split("\n"))
    return _BLANKS.sub("\n\n", text).strip()


def _first_heading(text: str) -> str | None:
    for line in text.split("\n"):
        line = line.strip()
        if line.startswith("#"):
            return line.lstrip("#").strip() or None
        if line:
            return line[:120]
    return None


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-16", "cp1252"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def parse_text(data: bytes, filename: str) -> ParsedDocument:
    text = normalize_whitespace(_decode(data))
    fmt = "markdown" if filename.lower().endswith((".md", ".markdown")) else "text"
    return ParsedDocument(text=text, title=_first_heading(text) or PurePath(filename).stem, fmt=fmt)


def parse_pdf(data: bytes, filename: str) -> ParsedDocument:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(data))
        pages = [(p.extract_text() or "") for p in reader.pages]
    except Exception as exc:  # pypdf raises many exception types for corrupt files
        raise ParseError(f"could not read PDF: {exc}") from exc
    meta = reader.metadata or {}
    # Keep page boundaries as form feeds so chunk metadata can carry page numbers.
    text = "\n\f\n".join(normalize_whitespace(p) for p in pages)
    title = (meta.get("/Title") or "").strip() or _first_heading(text) or PurePath(filename).stem
    metadata = {"pages": len(pages)}
    if meta.get("/Author"):
        metadata["author"] = str(meta.get("/Author"))
    if meta.get("/CreationDate"):
        metadata["created"] = str(meta.get("/CreationDate"))
    return ParsedDocument(text=text, title=title, fmt="pdf", metadata=metadata)


def parse_docx(data: bytes, filename: str) -> ParsedDocument:
    import docx

    try:
        d = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise ParseError(f"could not read DOCX: {exc}") from exc
    lines: list[str] = []
    for p in d.paragraphs:
        t = p.text.strip()
        if not t:
            continue
        style = (p.style.name or "").lower() if p.style is not None else ""
        if style.startswith("heading"):
            level = "".join(ch for ch in style if ch.isdigit()) or "1"
            t = "#" * int(level) + " " + t
        lines.append(t)
    for table in d.tables:
        for row in table.rows:
            lines.append(" | ".join(c.text.strip() for c in row.cells))
    text = normalize_whitespace("\n\n".join(lines))
    cp = d.core_properties
    title = (cp.title or "").strip() or _first_heading(text) or PurePath(filename).stem
    metadata = {}
    if cp.author:
        metadata["author"] = cp.author
    if cp.created:
        metadata["created"] = cp.created.isoformat()
    return ParsedDocument(text=text, title=title, fmt="docx", metadata=metadata)


def parse_html(data: bytes, filename: str) -> ParsedDocument:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(_decode(data), "html.parser")
    for tag in soup(["script", "style", "noscript", "nav", "footer"]):
        tag.decompose()
    for level in range(1, 7):
        for h in soup.find_all(f"h{level}"):
            h.insert_before("\n" + "#" * level + " ")
            h.insert_after("\n")
    for br in soup.find_all(["p", "li", "tr", "div", "br"]):
        br.insert_after("\n")
    text = normalize_whitespace(soup.get_text())
    title = (soup.title.string.strip() if soup.title and soup.title.string else None) or _first_heading(text)
    metadata = {}
    for m in soup.find_all("meta"):
        if m.get("name") in ("author", "date", "description") and m.get("content"):
            metadata[m["name"]] = m["content"]
    return ParsedDocument(text=text, title=title or PurePath(filename).stem, fmt="html", metadata=metadata)


def parse_json(data: bytes, filename: str) -> ParsedDocument:
    """JSON / JSONL: records are flattened to 'key: value' lines.

    If a record has a 'text'/'content'/'body' field it is used as the body.
    """
    raw = _decode(data)
    try:
        if filename.lower().endswith(".jsonl"):
            records = [json.loads(line) for line in raw.splitlines() if line.strip()]
        else:
            obj = json.loads(raw)
            records = obj if isinstance(obj, list) else [obj]
    except json.JSONDecodeError as exc:
        raise ParseError(f"invalid JSON: {exc}") from exc
    blocks = []
    title = None
    for rec in records:
        if isinstance(rec, dict):
            title = title or rec.get("title") or rec.get("name")
            body = rec.get("text") or rec.get("content") or rec.get("body")
            fields = [f"{k}: {v}" for k, v in rec.items() if k not in ("text", "content", "body") and not isinstance(v, (dict, list))]
            blocks.append("\n".join(fields + ([str(body)] if body else [])))
        else:
            blocks.append(str(rec))
    text = normalize_whitespace("\n\n".join(blocks))
    return ParsedDocument(text=text, title=str(title or PurePath(filename).stem), fmt="json", metadata={"records": len(records)})


PARSERS: dict[str, Callable[[bytes, str], ParsedDocument]] = {
    ".txt": parse_text,
    ".md": parse_text,
    ".markdown": parse_text,
    ".pdf": parse_pdf,
    ".docx": parse_docx,
    ".html": parse_html,
    ".htm": parse_html,
    ".json": parse_json,
    ".jsonl": parse_json,
}


def parse_file(data: bytes, filename: str) -> ParsedDocument:
    ext = PurePath(filename).suffix.lower()
    parser = PARSERS.get(ext)
    if parser is None:
        raise UnsupportedFormatError(f"unsupported file type '{ext}'. Supported: {', '.join(sorted(PARSERS))}")
    doc = parser(data, filename)
    if not doc.text.strip():
        raise ParseError("document contains no extractable text (scanned PDF without OCR?)")
    doc.metadata.setdefault("filename", PurePath(filename).name)
    return doc
